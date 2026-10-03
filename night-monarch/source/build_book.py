from pathlib import Path
import re, html, json, collections
import fitz
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_JUSTIFY
from reportlab.platypus import (BaseDocTemplate, PageTemplate, Frame, Paragraph, Spacer,
    PageBreak, NextPageTemplate, LongTable, TableStyle, Flowable)
from reportlab.platypus.tableofcontents import TableOfContents
from docx import Document
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
for name, file in [('Book','DejaVuSerif.ttf'),('Book-Bold','DejaVuSerif-Bold.ttf'),
                   ('Book-Italic','DejaVuSerif-Italic.ttf'),('Book-BoldItalic','DejaVuSerif-BoldItalic.ttf'),
                   ('Sans','DejaVuSans.ttf'),('Sans-Bold','DejaVuSans-Bold.ttf')]:
    pdfmetrics.registerFont(TTFont(name, '/usr/share/fonts/truetype/dejavu/'+file))
pdfmetrics.registerFontFamily('Book',normal='Book',bold='Book-Bold',italic='Book-Italic',boldItalic='Book-BoldItalic')
INK=colors.HexColor('#28241F'); RED=colors.HexColor('#7A3028'); GOLD=colors.HexColor('#AD8852')
PAPER=colors.HexColor('#FCF9F0'); LIGHT=colors.HexColor('#EEE6D4')
W,H=612,792; M=45; CW=(W-2*M-22)/2

events=json.loads((HERE/'book_content.json').read_text())
brand_names=['Lamia', 'Felyn', 'Dragon', 'Succubus', 'Angeloid (Fallen Angel)', 'Banshee', 'Jiangshi', 'Kitsune', 'Oni', 'Zinogre', 'Mermaid', 'Pixie', 'Gargoyle', 'Jorugumo', 'Specter', 'Celebremancer', 'Menreiki', 'Asura', 'Jaguar Sun', 'Rakshasa', 'Mummy Lord', 'Yuki-onna', 'Houri', 'Tressym', 'Manticore', 'Gryphon', 'Harlequin', 'Tengu', 'Unison Device']
styles={
 'p':ParagraphStyle('body',fontName='Book',fontSize=9.15,leading=12.6,textColor=INK,alignment=TA_JUSTIFY,spaceAfter=6,allowWidows=0,allowOrphans=0),
 'label':ParagraphStyle('label',fontName='Book-Italic',fontSize=8.5,leading=11.7,textColor=RED,spaceAfter=4,keepWithNext=True),
 'h2':ParagraphStyle('h2',fontName='Book-Bold',fontSize=12.6,leading=16,textColor=RED,spaceBefore=10,spaceAfter=5,keepWithNext=True),
 'h3':ParagraphStyle('h3',fontName='Book-Bold',fontSize=10.15,leading=13.6,textColor=RED,spaceBefore=7,spaceAfter=4,keepWithNext=True),
 'chapter':ParagraphStyle('chapter',fontName='Book-Bold',fontSize=24,leading=28,textColor=RED,spaceAfter=16,keepWithNext=True),
 'table':ParagraphStyle('cell',fontName='Book',fontSize=7.4,leading=10,textColor=INK),
 'thead':ParagraphStyle('thead',fontName='Sans-Bold',fontSize=7.1,leading=9.2,textColor=colors.white),
}

class Rule(Flowable):
    def __init__(self,width):Flowable.__init__(self);self.width=width;self.height=12;self.keepWithNext=True
    def draw(self):
        c=self.canv;c.setStrokeColor(GOLD);c.setLineWidth(.8);c.line(0,6,self.width,6)
        c.setFillColor(RED);c.rect(self.width/2-2,4,4,4,fill=1,stroke=0)

def background(c,doc):
    c.saveState();c.setFillColor(PAPER);c.rect(0,0,W,H,stroke=0,fill=1)
    c.setStrokeColor(GOLD);c.setLineWidth(.65)
    c.line(M, H-34,W-M,H-34);c.line(M,35,W-M,35)
    c.setFont('Sans',7);c.setFillColor(RED)
    c.drawString(M,H-26,'NIGHT MONARCH  •  STANDALONE CLASS')
    c.setFont('Book',8);c.drawCentredString(W/2,23,str(doc.page))
    if c._doctemplate.pageTemplate.id=='body':
        c.setStrokeColor(colors.HexColor('#DED4BF'));c.setLineWidth(.3);c.line(W/2,52,W/2,H-55)
    c.restoreState()

def page_end(c,doc):
    c.saveState();c.setFont('Sans',7);c.setFillColor(RED)
    c.drawRightString(W-M,H-26,getattr(doc,'chapter_name','Heir to the Night').upper());c.restoreState()

def cover(c,doc):
    c.saveState();c.setFillColor(colors.HexColor('#1F302C'));c.rect(0,0,W,H,fill=1,stroke=0)
    c.setStrokeColor(GOLD);c.setLineWidth(1.2);c.rect(24,24,W-48,H-48,fill=0)
    c.rect(31,31,W-62,H-62,fill=0)
    c.setFillColor(GOLD);c.circle(W/2,390,104,fill=1,stroke=0)
    c.setFillColor(colors.HexColor('#1F302C'));c.circle(W/2+25,407,99,fill=1,stroke=0)
    for x,y in [(W/2-80,285),(W/2,263),(W/2+80,285)]:
        c.setStrokeColor(GOLD);c.setFillColor(RED)
        p=c.beginPath();p.moveTo(x,y+15);p.lineTo(x+11,y);p.lineTo(x,y-15);p.lineTo(x-11,y);p.close()
        c.drawPath(p,fill=1,stroke=1)
    c.setFillColor(PAPER);c.setFont('Book-Bold',36);c.drawCentredString(W/2,614,'NIGHT MONARCH')
    c.setFillColor(GOLD);c.setFont('Sans-Bold',13);c.drawCentredString(W/2,574,'STANDALONE CLASS')
    c.setFillColor(PAPER);c.setFont('Book-Italic',17);c.drawCentredString(W/2,538,'Heir to the Night')
    c.setFont('Book',12);c.drawCentredString(W/2,216,'Soul Absorption • Bloodline Brands • Magatama')
    c.setFont('Book-Italic',10);c.drawCentredString(W/2,188,'A Standalone Soma Cruz–inspired class • Revised Master')
    c.setFillColor(GOLD);c.setFont('Sans',8);c.drawCentredString(W/2,91,'REPOSITORY RULES  /  PATHFINDER-INSPIRED LAYOUT EDITION')
    c.restoreState()

class BookDoc(BaseDocTemplate):
    def beforeDocument(self):self.chapter_name='Heir to the Night';self.heading_no=0
    def afterFlowable(self,f):
        if hasattr(f,'chapter_title'):self.chapter_name=f.chapter_title
        if hasattr(f,'toc_level'):
            self.heading_no+=1;key='section-'+str(self.heading_no)
            self.canv.bookmarkPage(key)
            self.canv.addOutlineEntry(f.getPlainText(),key,level=f.toc_level,closed=True)
            self.notify('TOCEntry',(f.toc_level,f.getPlainText(),self.page,key))

def table(spec,width):
    n=len(spec['headers'])
    if spec.get('wide'):
        weights=([.06,.13,.06,.06,.06,.23,.10,.09,.06,.11,.11] if n==11 else [.07,.07,.32,.09,.11,.11,.11,.12] if n==8 else [.07,.105,.062,.062,.062,.32,.125,.105,.055] if n==9 else
                 [.065,.065,.40,.105,.13,.13,.105] if n==7 else
                 [.07,.075,.385,.135,.175,.16] if n==6 else [.10,.22,.68] if n==3 else [.27,.73] if n==2 else [.07,.075,.49,.165,.20])
    elif n==3:weights=[.08,.25,.67]
    elif spec['name']=='Ofuda Modifications':weights=[.29,.71]
    else:weights=[.34,.66]
    ws=[width*v/sum(weights) for v in weights]
    data=[[Paragraph(html.escape(v),styles['thead']) for v in spec['headers']]]
    for row in spec['rows']:
        data.append([Paragraph(html.escape(v),styles['table']) for v in row])
    t=LongTable(data,colWidths=ws,repeatRows=1,hAlign='LEFT',spaceAfter=10)
    t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),RED),('ROWBACKGROUNDS',(0,1),(-1,-1),[PAPER,LIGHT]),
       ('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5),
       ('TOPPADDING',(0,0),(-1,-1),4),('BOTTOMPADDING',(0,0),(-1,-1),4),
       ('LINEBELOW',(0,0),(-1,0),1,GOLD),('LINEBELOW',(0,-1),(-1,-1),.5,GOLD)]))
    return t

def make_pdf():
    doc=BookDoc(str(HERE/'Night-Monarch.pdf'),pagesize=(W,H),leftMargin=M,rightMargin=M,topMargin=50,bottomMargin=50,
        title='Night Monarch — Revised Master',author='Night Monarch source / current layout edition')
    doc.addPageTemplates([
       PageTemplate(id='cover',frames=[Frame(M,50,W-2*M,H-100,id='coverframe',leftPadding=0,rightPadding=0,topPadding=0,bottomPadding=0)],onPage=cover),
       PageTemplate(id='wide',frames=[Frame(M,50,W-2*M,H-100,id='wideframe',leftPadding=0,rightPadding=0,topPadding=0,bottomPadding=0)],onPage=background,onPageEnd=page_end),
       PageTemplate(id='body',frames=[Frame(M,50,CW,H-100,id='left',leftPadding=0,rightPadding=0,topPadding=0,bottomPadding=0),Frame(M+CW+22,50,CW,H-100,id='right',leftPadding=0,rightPadding=0,topPadding=0,bottomPadding=0)],onPage=background,onPageEnd=page_end),
    ])
    story=[Spacer(1,1),NextPageTemplate('wide'),PageBreak(),Paragraph('Contents',styles['chapter']),Rule(W-2*M)]
    toc=TableOfContents();toc.levelStyles=[ParagraphStyle('toc0',fontName='Book-Bold',fontSize=10,leading=16,textColor=RED,spaceBefore=4),ParagraphStyle('toc1',fontName='Book',fontSize=8.5,leading=12,textColor=INK,leftIndent=12)]
    story.append(toc)
    bigs=[e for e in events if e['kind']=='table']
    for e in bigs:
        story.extend([NextPageTemplate('wide'),PageBreak()])
        p=Paragraph(e['spec']['name'],styles['h2']);p.toc_level=0;p.chapter_title=e['spec']['name']
        story.extend([p,Rule(W-2*M),table(e['spec'],W-2*M)])
    story.extend([NextPageTemplate('body'),PageBreak()])
    for e in events:
        kind=e['kind']
        if kind=='table':continue
        if kind=='chapter':
            p=Paragraph(html.escape(e['text']),styles['h2']);p.toc_level=0;p.chapter_title=e['text'];story.extend([Spacer(1,10),p,Rule(CW)])
        elif kind=='h':
            p=Paragraph(html.escape(e['text']),styles['h'+str(e['level'])])
            if e['level']==2 and e['text'] in brand_names:p.toc_level=1
            story.append(p)
        else:story.append(Paragraph(e.get('rich',html.escape(e['text'])),styles[kind]))
    doc.multiBuild(story)

def make_docx():
    doc=Document();section=doc.sections[0]
    section.top_margin=section.bottom_margin=Inches(.7)
    section.left_margin=section.right_margin=Inches(.7)
    normal=doc.styles['Normal'];normal.font.name='Georgia';normal.font.size=Pt(10)
    normal.paragraph_format.space_after=Pt(6)
    for style in ['Title','Heading 1','Heading 2','Heading 3']:
        doc.styles[style].font.name='Georgia';doc.styles[style].font.color.rgb=RGBColor.from_string('7A3028')
    doc.add_heading('NIGHT MONARCH — STANDALONE CLASS',0);doc.add_paragraph('Heir to the Night')
    doc.add_paragraph('Standalone class; revised-master rules with recovered source material. Unfinished donor interfaces are identified in the revision ledger.')
    for e in events:
        kind=e['kind']
        if kind=='chapter':doc.add_page_break();doc.add_heading(e['text'],1)
        elif kind=='h':doc.add_heading(e['text'],min(e['level']+1,3))
        elif kind=='table':
            spec=e['spec'];doc.add_heading(spec['name'],3)
            t=doc.add_table(rows=1,cols=len(spec['headers']));t.style='Light Shading Accent 1'
            for cell,v in zip(t.rows[0].cells,spec['headers']):cell.text=v
            for row in spec['rows']:
                for cell,v in zip(t.add_row().cells,row):cell.text=v
        else:doc.add_paragraph(e['text'])
    doc.save(HERE/'Night-Monarch-Editable.docx')

make_pdf();make_docx()
print('Created PDF and editable DOCX:',HERE)
print('Chapters:',[e['text'] for e in events if e['kind']=='chapter'])
print('Tables:',[(e['spec']['name'],len(e['spec']['rows'])) for e in events if e['kind']=='table'])
