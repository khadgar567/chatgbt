#!/usr/bin/env python3
"""Build the audited FX framework as a publication-quality PDF.

The DOCX remains the content authority.  This script deliberately performs no
market-data refresh: it only maps the source paragraphs, runs, tables, and one
embedded exhibit into a stable PDF layout.
"""

from __future__ import annotations

import html
import io
import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph
from PIL import Image as PILImage
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    Image,
    KeepTogether,
    LongTable,
    NextPageTemplate,
    PageBreak,
    PageTemplate,
    Paragraph,
    Spacer,
    TableStyle,
)
from reportlab.platypus.tableofcontents import TableOfContents


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "FX_Framework_v9_2_Audited.docx"
OUTPUT = ROOT / "FX_Framework_v9_2_Rebuilt.pdf"

INK = colors.HexColor("#152536")
NAVY = colors.HexColor("#173B57")
TEAL = colors.HexColor("#147D7E")
GOLD = colors.HexColor("#C69B47")
PALE = colors.HexColor("#EAF2F4")
PALE_GOLD = colors.HexColor("#F7F1E5")
GRID = colors.HexColor("#A8B7BF")
MUTED = colors.HexColor("#5D6B73")


def register_fonts() -> None:
    font_dir = Path("/usr/share/fonts/truetype/dejavu")
    for name, filename in {
        "DejaVu": "DejaVuSans.ttf",
        "DejaVu-Bold": "DejaVuSans-Bold.ttf",
        # This minimal container ships Sans regular/bold and Mono italics.
        # The latter are used only for the source's sparse emphasis runs.
        "DejaVu-Oblique": "DejaVuSansMono-Oblique.ttf",
        "DejaVu-BoldOblique": "DejaVuSansMono-BoldOblique.ttf",
        "DejaVuMono": "DejaVuSansMono.ttf",
    }.items():
        pdfmetrics.registerFont(TTFont(name, str(font_dir / filename)))
    pdfmetrics.registerFontFamily("DejaVu", normal="DejaVu", bold="DejaVu-Bold",
                                  italic="DejaVu-Oblique", boldItalic="DejaVu-BoldOblique")


def iter_blocks(document):
    """Yield source paragraphs and tables in their actual document order."""
    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield DocxParagraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield DocxTable(child, document)


def clean_text(text: str) -> str:
    return text.replace("\u00a0", " ").replace("\u2011", "-").strip()


def rich_paragraph_text(paragraph: DocxParagraph) -> str:
    """Preserve emphasis while producing safe ReportLab mini-markup."""
    pieces: list[str] = []
    # iter_inner_content includes hyperlink children, while paragraph.runs does
    # not. Omitting those children would silently drop the audit source URLs.
    for item in paragraph.iter_inner_content():
        is_link = item.__class__.__name__ == "Hyperlink"
        runs = item.runs if is_link else [item]
        for run in runs:
            text = html.escape(run.text).replace("\n", "<br/>").replace("\t", "&nbsp;&nbsp;&nbsp;")
            if not text:
                continue
            if run.bold and run.italic:
                text = f"<b><i>{text}</i></b>"
            elif run.bold:
                text = f"<b>{text}</b>"
            elif run.italic:
                text = f"<i>{text}</i>"
            if run.underline:
                text = f"<u>{text}</u>"
            if is_link and item.url:
                text = f'<link href="{html.escape(item.url)}" color="#147D7E">{text}</link>'
            pieces.append(text)
    return "".join(pieces) or html.escape(clean_text(paragraph.text))


class FrameworkDoc(BaseDocTemplate):
    def __init__(self, filename: str, **kwargs):
        super().__init__(filename, **kwargs)
        self.heading_no = 0

    def beforeDocument(self):
        # multiBuild lays the story out more than once while resolving the TOC.
        self.heading_no = 0

    def afterFlowable(self, flowable):
        if not isinstance(flowable, Paragraph):
            return
        level = getattr(flowable, "toc_level", None)
        if level is None:
            return
        self.heading_no += 1
        key = f"heading-{self.heading_no}"
        self.canv.bookmarkPage(key)
        self.canv.addOutlineEntry(flowable.getPlainText(), key, level=level, closed=False)
        self.notify("TOCEntry", (level, flowable.getPlainText(), self.page, key))


def page_header_footer(canvas, doc):
    canvas.saveState()
    width, height = A4
    canvas.setStrokeColor(GOLD)
    canvas.setLineWidth(0.7)
    canvas.line(doc.leftMargin, height - 16 * mm, width - doc.rightMargin, height - 16 * mm)
    canvas.setFont("DejaVu-Bold", 7.5)
    canvas.setFillColor(NAVY)
    canvas.drawString(doc.leftMargin, height - 12.5 * mm, "SYSTEMIC FX COMPRESSION & GOLD PROXY REVALUATION")
    canvas.setFont("DejaVu", 7.5)
    canvas.setFillColor(MUTED)
    canvas.drawRightString(width - doc.rightMargin, height - 12.5 * mm, "v9.2  •  AUDITED WORKING MASTER")
    canvas.setStrokeColor(colors.HexColor("#D8E0E4"))
    canvas.line(doc.leftMargin, 14 * mm, width - doc.rightMargin, 14 * mm)
    canvas.setFont("DejaVu", 7.3)
    canvas.setFillColor(MUTED)
    canvas.drawString(doc.leftMargin, 9.5 * mm, "CONFIDENTIAL — INTERNAL WORKING MASTER")
    canvas.drawRightString(width - doc.rightMargin, 9.5 * mm, f"PAGE {doc.page}")
    canvas.restoreState()


def cover(canvas, doc):
    canvas.saveState()
    width, height = A4
    canvas.setFillColor(NAVY)
    canvas.rect(0, 0, width, height, fill=1, stroke=0)
    canvas.setFillColor(TEAL)
    canvas.rect(0, 0, 16 * mm, height, fill=1, stroke=0)
    canvas.setFillColor(GOLD)
    canvas.rect(23 * mm, height - 62 * mm, 34 * mm, 2.4 * mm, fill=1, stroke=0)
    canvas.setFont("DejaVu-Bold", 12)
    canvas.setFillColor(colors.HexColor("#9EC7C7"))
    canvas.drawString(23 * mm, height - 45 * mm, "AUDITED FRAMEWORK  /  VERSION 9.2")
    canvas.setFont("DejaVu-Bold", 27)
    canvas.setFillColor(colors.white)
    title = ["SYSTEMIC FX", "COMPRESSION &", "GOLD PROXY", "REVALUATION", "FRAMEWORK"]
    y = height - 83 * mm
    for line in title:
        canvas.drawString(23 * mm, y, line)
        y -= 13 * mm
    canvas.setFont("DejaVu", 13)
    canvas.setFillColor(colors.HexColor("#D7E5E8"))
    canvas.drawString(23 * mm, y - 4 * mm, "Sentinel Audit & Author-Intent Reconciliation")
    canvas.setStrokeColor(colors.HexColor("#507187"))
    canvas.line(23 * mm, 62 * mm, width - 23 * mm, 62 * mm)
    canvas.setFont("DejaVu-Bold", 9)
    canvas.setFillColor(GOLD)
    canvas.drawString(23 * mm, 52 * mm, "AUTHOR")
    canvas.drawString(80 * mm, 52 * mm, "REVISION")
    canvas.drawString(139 * mm, 52 * mm, "HISTORICAL SNAPSHOT")
    canvas.setFont("DejaVu", 10)
    canvas.setFillColor(colors.white)
    canvas.drawString(23 * mm, 45 * mm, "Kaan")
    canvas.drawString(80 * mm, 45 * mm, "8 September 2026")
    canvas.drawString(139 * mm, 45 * mm, "4 August 2026")
    canvas.setFont("DejaVu-Bold", 8)
    canvas.setFillColor(colors.HexColor("#AFC5CE"))
    canvas.drawString(23 * mm, 24 * mm, "CONFIDENTIAL — INTERNAL WORKING MASTER")
    canvas.restoreState()


def styles():
    base = getSampleStyleSheet()
    return {
        "body": ParagraphStyle("Body", parent=base["BodyText"], fontName="DejaVu", fontSize=8.65,
                               leading=12.3, textColor=INK, spaceAfter=5.2, alignment=TA_LEFT,
                               splitLongWords=False, allowWidows=0, allowOrphans=0),
        "h1": ParagraphStyle("H1", fontName="DejaVu-Bold", fontSize=16, leading=19,
                             textColor=NAVY, spaceBefore=13, spaceAfter=7, keepWithNext=True),
        "h2": ParagraphStyle("H2", fontName="DejaVu-Bold", fontSize=11.3, leading=14,
                             textColor=TEAL, spaceBefore=9, spaceAfter=5, keepWithNext=True),
        "h3": ParagraphStyle("H3", fontName="DejaVu-Bold", fontSize=9.5, leading=12,
                             textColor=NAVY, spaceBefore=7, spaceAfter=4, keepWithNext=True),
        "toc_title": ParagraphStyle("TOCTitle", fontName="DejaVu-Bold", fontSize=19,
                                    leading=23, textColor=NAVY, spaceAfter=10),
        "table": ParagraphStyle("TableCell", fontName="DejaVu", fontSize=6.9, leading=8.7,
                                textColor=INK, splitLongWords=True),
        "table_head": ParagraphStyle("TableHead", fontName="DejaVu-Bold", fontSize=7,
                                     leading=8.8, textColor=colors.white),
        "caption": ParagraphStyle("Caption", fontName="DejaVu-Oblique", fontSize=7.2,
                                  leading=9, textColor=MUTED, spaceAfter=4),
    }


def table_widths(table: DocxTable, available: float) -> list[float]:
    cols = len(table.columns)
    scores = []
    for c in range(cols):
        lengths = [max(6, min(70, len(clean_text(row.cells[c].text)))) for row in table.rows]
        scores.append(max(12, sum(lengths) / max(1, len(lengths))))
    floor = available * (0.105 if cols >= 5 else 0.15 if cols >= 4 else 0.18)
    raw = [max(floor, available * s / sum(scores)) for s in scores]
    scale = available / sum(raw)
    return [x * scale for x in raw]


def make_table(source: DocxTable, sty, available: float):
    cols = len(source.columns)
    cell_style = sty["table"]
    head_style = sty["table_head"]
    data = []
    for r, row in enumerate(source.rows):
        converted = []
        for cell in row.cells:
            value = html.escape(clean_text(cell.text)).replace("\n", "<br/>") or "—"
            converted.append(Paragraph(value, head_style if r == 0 else cell_style))
        data.append(converted)
    t = LongTable(data, colWidths=table_widths(source, available), repeatRows=1,
                  hAlign="LEFT", splitByRow=1, spaceBefore=4, spaceAfter=9)
    commands = [
        ("BACKGROUND", (0, 0), (-1, 0), NAVY),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.35, GRID),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]
    for r in range(1, len(data)):
        commands.append(("BACKGROUND", (0, r), (-1, r), colors.white if r % 2 else PALE))
    if cols >= 6:
        commands.extend([("LEFTPADDING", (0, 0), (-1, -1), 2.5),
                         ("RIGHTPADDING", (0, 0), (-1, -1), 2.5)])
    t.setStyle(TableStyle(commands))
    return t


def extract_image(paragraph: DocxParagraph, document: Document) -> bytes | None:
    blips = paragraph._p.xpath(".//a:blip")
    if not blips:
        return None
    rid = blips[0].get(qn("r:embed"))
    return document.part.related_parts[rid].blob


def image_flowable(blob: bytes, available_width: float):
    im = PILImage.open(io.BytesIO(blob))
    width, height = im.size
    scale = min(available_width / width, 115 * mm / height)
    result = Image(io.BytesIO(blob), width=width * scale, height=height * scale)
    result.hAlign = "CENTER"
    return result


def classify_body(text: str, paragraph: DocxParagraph) -> str:
    style = paragraph.style.name
    if text == "Abstract":
        return "h1"
    if style == "FX Heading 1":
        return "h1"
    if style == "FX Heading 2":
        return "h2"
    # The source has several numbered subheads stored as body paragraphs.
    if re.match(r"^(?:\d+(?:\.\d+){1,3}[A-Z]?|Appendix [A-Z](?:\.\d+)?)\s*[.\-—:]", text):
        return "h3"
    return "body"


def build() -> None:
    register_fonts()
    source = Document(SOURCE)
    sty = styles()
    margin = 19 * mm
    doc = FrameworkDoc(str(OUTPUT), pagesize=A4, leftMargin=margin, rightMargin=margin,
                       topMargin=21 * mm, bottomMargin=18 * mm,
                       title="Systemic FX Compression & Gold Proxy Revaluation Framework v9.2",
                       author="Kaan", subject="Sentinel Audit & Author-Intent Reconciliation")
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="main")
    doc.addPageTemplates([
        PageTemplate(id="cover", frames=frame, onPage=cover),
        PageTemplate(id="body", frames=frame, onPage=page_header_footer),
    ])

    story = [Spacer(1, 1), NextPageTemplate("body"), PageBreak()]
    story.append(Paragraph("Contents", sty["toc_title"]))
    toc = TableOfContents()
    toc.levelStyles = [
        ParagraphStyle("TOC1", fontName="DejaVu-Bold", fontSize=8.5, leading=12,
                       leftIndent=0, firstLineIndent=0, textColor=NAVY, spaceBefore=2),
        ParagraphStyle("TOC2", fontName="DejaVu", fontSize=7.7, leading=10.5,
                       leftIndent=9 * mm, firstLineIndent=0, textColor=INK),
        ParagraphStyle("TOC3", fontName="DejaVu", fontSize=7.2, leading=9.6,
                       leftIndent=17 * mm, firstLineIndent=0, textColor=MUTED),
    ]
    story.extend([toc, PageBreak()])

    title_paragraphs_skipped = 0
    table_number = 0
    for block in iter_blocks(source):
        if isinstance(block, DocxParagraph):
            raw = clean_text(block.text)
            blob = extract_image(block, source)
            if blob:
                story.extend([Spacer(1, 4), image_flowable(blob, doc.width),
                              Paragraph("Genesis exhibit preserved from the authoritative source document.", sty["caption"]),
                              Spacer(1, 5)])
                continue
            if not raw:
                continue
            if title_paragraphs_skipped < 3:
                title_paragraphs_skipped += 1
                continue
            kind = classify_body(raw, block)
            p = Paragraph(rich_paragraph_text(block), sty[kind])
            if kind.startswith("h"):
                p.toc_level = {"h1": 0, "h2": 1, "h3": 2}[kind]
            # Keep short lead-ins with the paragraph that follows where possible.
            if kind == "body" and len(raw) < 95 and (raw.endswith(":") or raw.startswith("Core thesis")):
                p.keepWithNext = True
            story.append(p)
        else:
            table_number += 1
            story.append(make_table(block, sty, doc.width))

    doc.multiBuild(story)
    print(f"Built {OUTPUT} from {len(source.paragraphs)} paragraphs and {table_number} tables")


if __name__ == "__main__":
    build()
